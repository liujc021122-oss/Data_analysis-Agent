import base64
import inspect
from zipfile import ZipFile
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

from data_analysis_agent.reports.word import generate_word_report
from utils.word_report_generator import generate_word_report as compatibility_generate_word_report


TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def write_tiny_png(path: Path) -> None:
    path.write_bytes(TINY_PNG)


def test_compatibility_export_keeps_generate_word_report_signature():
    assert compatibility_generate_word_report is generate_word_report
    signature = inspect.signature(generate_word_report)
    assert list(signature.parameters) == [
        "markdown_content",
        "output_path",
        "session_output_dir",
        "figures",
    ]
    assert signature.parameters["figures"].default is None


def test_generates_docx_with_markdown_structure_and_embedded_image(tmp_path):
    session_dir = tmp_path / "session"
    session_dir.mkdir()
    image_path = session_dir / "趋势图.png"
    write_tiny_png(image_path)
    output_path = session_dir / "最终分析报告.docx"

    markdown = """# 数据分析报告

这是报告摘要。

## 关键发现

- 第一项发现
- 第二项发现

### 图表分析

![趋势图](./趋势图.png)

**重点**和*说明*。
"""

    result = generate_word_report(
        markdown_content=markdown,
        output_path=output_path,
        session_output_dir=session_dir,
    )

    assert result == str(output_path)
    assert output_path.exists()

    document = Document(output_path)
    paragraphs = document.paragraphs
    paragraph_text = "\n".join(paragraph.text for paragraph in paragraphs)

    assert "数据分析报告" in paragraph_text
    assert "这是报告摘要。" in paragraph_text
    assert "第一项发现" in paragraph_text
    assert "第二项发现" in paragraph_text
    assert any(paragraph.style.name == "Heading 1" for paragraph in paragraphs)
    assert any(paragraph.style.name == "List Bullet" for paragraph in paragraphs)
    assert len(document.inline_shapes) == 1
    assert any(run.bold for paragraph in paragraphs for run in paragraph.runs)
    assert any(run.italic for paragraph in paragraphs for run in paragraph.runs)


def test_missing_or_outside_image_is_reported_without_breaking_docx(tmp_path):
    session_dir = tmp_path / "session"
    session_dir.mkdir()
    outside_image = tmp_path / "outside.png"
    write_tiny_png(outside_image)
    output_path = session_dir / "最终分析报告.docx"

    markdown = """# 报告

![缺失图片](./missing.png)
![越界图片](../outside.png)
"""

    generate_word_report(
        markdown_content=markdown,
        output_path=output_path,
        session_output_dir=session_dir,
    )

    document = Document(output_path)
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "图片不可用" in paragraph_text
    assert len(document.inline_shapes) == 0


def test_heading_runs_inherit_heading_style_tokens(tmp_path):
    session_dir = tmp_path / "session"
    session_dir.mkdir()
    output_path = session_dir / "最终分析报告.docx"

    generate_word_report(
        markdown_content="# 标题\n\n正文。",
        output_path=output_path,
        session_output_dir=session_dir,
    )

    document = Document(output_path)
    heading = next(paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 1")

    assert heading.runs[0].font.size is None
    assert heading.runs[0].font.color.rgb is None


def test_collected_figure_path_can_resolve_a_different_display_filename(tmp_path):
    session_dir = tmp_path / "session"
    session_dir.mkdir()
    image_path = session_dir / "generated" / "actual-name.png"
    image_path.parent.mkdir()
    write_tiny_png(image_path)
    output_path = session_dir / "最终分析报告.docx"

    generate_word_report(
        markdown_content="# 报告\n\n![趋势图](./display-name.png)",
        output_path=output_path,
        session_output_dir=session_dir,
        figures=[
            {
                "filename": "display-name.png",
                "file_path": str(image_path),
            }
        ],
    )

    document = Document(output_path)

    assert len(document.inline_shapes) == 1


def test_renders_section_summary_and_gray_italic_analysis_points(tmp_path):
    session_dir = tmp_path / "session"
    session_dir.mkdir()
    output_path = session_dir / "最终分析报告.docx"

    markdown = """# 数据分析报告

## 营收分析

【部分总结】各门店收入集中度较高，整体经营仍有改善空间。
【分析要点】
- 前四名门店贡献主要收入。
- 疫情期间收入出现明显波动。
"""

    generate_word_report(
        markdown_content=markdown,
        output_path=output_path,
        session_output_dir=session_dir,
    )

    document = Document(output_path)
    summary = next(
        paragraph for paragraph in document.paragraphs
        if "各门店收入集中度较高" in paragraph.text
    )
    points = [
        paragraph for paragraph in document.paragraphs
        if paragraph.text in {"前四名门店贡献主要收入。", "疫情期间收入出现明显波动。"}
    ]

    assert summary.style.name == "Analysis Summary"
    assert summary.runs[0].font.color.rgb == RGBColor(0xFF, 0x72, 0x00)
    assert summary.runs[0].bold is True
    assert len(points) == 2
    assert all(paragraph.style.name == "List Bullet" for paragraph in points)
    assert all(run.font.color.rgb == RGBColor(0x8C, 0x8C, 0x8C) for paragraph in points for run in paragraph.runs)
    assert all(run.italic is True for paragraph in points for run in paragraph.runs)


def test_uses_a4_page_and_unified_heading_font(tmp_path):
    session_dir = tmp_path / "session"
    session_dir.mkdir()
    output_path = session_dir / "最终分析报告.docx"

    generate_word_report(
        markdown_content="# 报告\n\n## 分析\n\n### 细分\n\n正文。",
        output_path=output_path,
        session_output_dir=session_dir,
    )

    document = Document(output_path)
    section = document.sections[0]

    assert round(section.page_width.inches, 2) == 8.27
    assert round(section.page_height.inches, 2) == 11.69
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        assert style.font.name == "Microsoft YaHei"
        r_fonts = style._element.rPr.rFonts
        assert r_fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii") == "Microsoft YaHei"
        assert r_fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia") == "Microsoft YaHei"
    assert all(
        paragraph.runs[0].font.name == "Microsoft YaHei"
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading") and paragraph.runs
    )


def test_embedded_image_has_border_and_outer_shadow(tmp_path):
    session_dir = tmp_path / "session"
    session_dir.mkdir()
    image_path = session_dir / "趋势图.png"
    write_tiny_png(image_path)
    output_path = session_dir / "最终分析报告.docx"

    generate_word_report(
        markdown_content="# 报告\n\n![趋势图](./趋势图.png)",
        output_path=output_path,
        session_output_dir=session_dir,
    )

    with ZipFile(output_path) as archive:
        xml = archive.read("word/document.xml")

    assert b"<a:ln" in xml
    assert b"w=\"12700\"" in xml
    assert b"<a:srgbClr val=\"D9D9D9\"" in xml
    assert b"<a:outerShdw" in xml
    assert b"<a:alpha val=\"25000\"" in xml
