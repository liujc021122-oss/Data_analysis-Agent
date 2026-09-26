"""Convert the generated Markdown report into a self-contained Word document."""

import re
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Optional, Union
from urllib.parse import unquote

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from data_analysis_agent.domain.enums import ReportFormat

from .models import ReportDocument


PathLike = Union[str, Path]

_HEADING_PATTERN = re.compile(r"^(#{1,3})\s+(.+?)\s*#*\s*$")
_BULLET_PATTERN = re.compile(r"^\s*[-*+]\s+(.+)$")
_NUMBER_PATTERN = re.compile(r"^\s*\d+[.)]\s+(.+)$")
_IMAGE_PATTERN = re.compile(r"!\[([^\]]*)\]\((.+)\)")
_SUMMARY_PATTERN = re.compile(
    r"^\s*(?:\*\*)?【部分总结】(?:\*\*)?\s*(?:(?:：|:)\s*)?(.*?)\s*$"
)
_POINTS_HEADER_PATTERN = re.compile(
    r"^\s*(?:\*\*)?【分析要点】(?:\*\*)?\s*(?:(?:：|:)\s*)?$"
)
_INLINE_PATTERN = re.compile(
    r"(\*\*[^*\n]+\*\*|__[^_\n]+__|\*[^*\n]+\*|_[^_\n]+_)"
)

_BLACK = RGBColor(0x00, 0x00, 0x00)
_HEADING_GRAY = RGBColor(0x40, 0x40, 0x40)
_ORANGE = RGBColor(0xFF, 0x72, 0x00)
_POINT_GRAY = RGBColor(0x8C, 0x8C, 0x8C)
_MUTED = RGBColor(0x66, 0x66, 0x66)
_BODY_FONT = "Calibri"
_HEADING_FONT = "Microsoft YaHei"


def _set_font_mapping(target: Any, ascii_font: str, east_asia_font: str) -> None:
    """Set Word's four font slots so Chinese text has an explicit fallback."""
    r_pr = target._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)
    r_fonts.set(qn("w:cs"), ascii_font)
    r_fonts.set(qn("w:eastAsia"), east_asia_font)


def _set_style_font(
    style: Any,
    size: float,
    color: RGBColor = _BLACK,
    bold: bool = False,
    font_name: str = _BODY_FONT,
    east_asia_font: str = "Microsoft YaHei",
) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    _set_font_mapping(style, font_name, east_asia_font)


def _set_run_font(run: Any, size: float = 11, color: RGBColor = _BLACK) -> None:
    run.font.name = _BODY_FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    _set_font_mapping(run, _BODY_FONT, "Microsoft YaHei")


def _is_within(path: Path, directory: Path) -> bool:
    try:
        path.relative_to(directory)
        return True
    except ValueError:
        return False


class WordReportGenerator:
    """Render the Markdown subset emitted by the analysis agent into DOCX."""

    def __init__(
        self,
        session_output_dir: PathLike,
        figures: Optional[Iterable[Dict[str, Any]]] = None,
    ) -> None:
        self.session_output_dir = Path(session_output_dir).resolve()
        self.figures = list(figures or [])
        self.document = Document()
        self._figure_paths = self._index_figures(self.figures)
        self._configure_document()

    def generate(self, markdown_content: str, output_path: PathLike) -> str:
        """Convert Markdown content and save the resulting document."""
        self._render_markdown(markdown_content or "")
        destination = Path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(destination))
        return str(destination)

    def _configure_document(self) -> None:
        section = self.document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)

        styles = self.document.styles
        normal = styles["Normal"]
        _set_style_font(normal, 11)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.10

        title = styles["Title"]
        _set_style_font(
            title,
            23,
            _HEADING_GRAY,
            bold=True,
            font_name=_HEADING_FONT,
            east_asia_font=_HEADING_FONT,
        )
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(8)

        subtitle = styles["Subtitle"]
        _set_style_font(
            subtitle,
            13,
            _MUTED,
            font_name=_HEADING_FONT,
            east_asia_font=_HEADING_FONT,
        )
        subtitle.paragraph_format.space_before = Pt(0)
        subtitle.paragraph_format.space_after = Pt(12)

        for name, size, before, after in (
            ("Heading 1", 18, 16, 8),
            ("Heading 2", 15, 12, 6),
            ("Heading 3", 12, 8, 4),
        ):
            style = styles[name]
            _set_style_font(
                style,
                size,
                _HEADING_GRAY,
                bold=True,
                font_name=_HEADING_FONT,
                east_asia_font=_HEADING_FONT,
            )
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True

        summary = styles.add_style("Analysis Summary", WD_STYLE_TYPE.PARAGRAPH)
        _set_style_font(
            summary,
            12,
            _ORANGE,
            bold=True,
            font_name=_HEADING_FONT,
            east_asia_font=_HEADING_FONT,
        )
        summary.paragraph_format.space_before = Pt(4)
        summary.paragraph_format.space_after = Pt(5)
        summary.paragraph_format.line_spacing = 1.10
        summary.paragraph_format.keep_with_next = True

        figure = styles.add_style("Figure", WD_STYLE_TYPE.PARAGRAPH)
        _set_style_font(figure, 10, _BLACK)
        figure.paragraph_format.space_before = Pt(5)
        figure.paragraph_format.space_after = Pt(2)
        figure.paragraph_format.line_spacing = 1.0
        figure.paragraph_format.keep_with_next = True

        for name in ("List Bullet", "List Number"):
            style = styles[name]
            _set_style_font(style, 11)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(8)
            style.paragraph_format.line_spacing = 1.167

        caption = styles["Caption"]
        _set_style_font(caption, 9, _MUTED)
        caption.font.italic = True
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(8)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(0)
        footer.paragraph_format.space_after = Pt(0)
        footer_run = footer.add_run("数据分析报告 | 自动生成")
        _set_run_font(footer_run, 9, _MUTED)

    @staticmethod
    def _index_figures(figures: Iterable[Dict[str, Any]]) -> Dict[str, Path]:
        indexed: Dict[str, Path] = {}
        for figure in figures:
            filename = figure.get("filename")
            file_path = figure.get("file_path")
            if filename and file_path:
                indexed[str(filename)] = Path(file_path)
                indexed[Path(str(filename)).name] = Path(file_path)
        return indexed

    def _render_markdown(self, markdown_content: str) -> None:
        lines = markdown_content.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        paragraph_lines = []
        in_code_block = False
        in_analysis_points = False
        pending_summary = False

        def flush_paragraph() -> None:
            if not paragraph_lines:
                return
            paragraph = self.document.add_paragraph(style="Normal")
            self._add_inline_runs(paragraph, " ".join(paragraph_lines).strip())
            paragraph_lines.clear()

        for raw_line in lines:
            line = raw_line.rstrip()
            stripped = line.strip()

            if stripped.startswith("```"):
                flush_paragraph()
                in_code_block = not in_code_block
                in_analysis_points = False
                pending_summary = False
                continue

            if in_code_block:
                paragraph = self.document.add_paragraph(style="No Spacing")
                run = paragraph.add_run(line)
                _set_run_font(run, 9, _MUTED)
                run.font.name = "Consolas"
                continue

            if not stripped:
                flush_paragraph()
                continue

            summary_match = _SUMMARY_PATTERN.match(stripped)
            if summary_match:
                flush_paragraph()
                in_analysis_points = False
                summary_text = summary_match.group(1).strip()
                if summary_text:
                    self._add_summary(summary_text)
                    pending_summary = False
                else:
                    pending_summary = True
                continue

            if _POINTS_HEADER_PATTERN.match(stripped):
                flush_paragraph()
                pending_summary = False
                in_analysis_points = True
                continue

            if pending_summary:
                if not (
                    _HEADING_PATTERN.match(stripped)
                    or _IMAGE_PATTERN.fullmatch(stripped)
                    or _BULLET_PATTERN.match(line)
                    or _NUMBER_PATTERN.match(line)
                ):
                    self._add_summary(stripped)
                    pending_summary = False
                    continue
                pending_summary = False

            image_match = _IMAGE_PATTERN.fullmatch(stripped)
            if image_match:
                flush_paragraph()
                in_analysis_points = False
                self._add_image(image_match.group(1), image_match.group(2))
                continue

            heading_match = _HEADING_PATTERN.match(stripped)
            if heading_match:
                flush_paragraph()
                in_analysis_points = False
                level = len(heading_match.group(1))
                paragraph = self.document.add_paragraph(style=f"Heading {level}")
                self._add_inline_runs(
                    paragraph,
                    heading_match.group(2).strip(),
                    font_name=_HEADING_FONT,
                    east_asia_font=_HEADING_FONT,
                )
                continue

            bullet_match = _BULLET_PATTERN.match(line)
            if bullet_match:
                flush_paragraph()
                paragraph = self.document.add_paragraph(style="List Bullet")
                if in_analysis_points:
                    self._add_inline_runs(
                        paragraph,
                        bullet_match.group(1).strip(),
                        font_color=_POINT_GRAY,
                        font_italic=True,
                    )
                else:
                    self._add_inline_runs(paragraph, bullet_match.group(1).strip())
                continue

            number_match = _NUMBER_PATTERN.match(line)
            if number_match:
                flush_paragraph()
                in_analysis_points = False
                paragraph = self.document.add_paragraph(style="List Number")
                self._add_inline_runs(paragraph, number_match.group(1).strip())
                continue

            if re.fullmatch(r"\s{0,3}([-*_])(?:\s*\1){2,}\s*", line):
                flush_paragraph()
                in_analysis_points = False
                paragraph = self.document.add_paragraph(style="Normal")
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                run = paragraph.add_run("________________________________")
                _set_run_font(run, 9, _MUTED)
                continue

            in_analysis_points = False
            paragraph_lines.append(stripped)

        flush_paragraph()

    def _add_summary(self, text: str) -> None:
        paragraph = self.document.add_paragraph(style="Analysis Summary")
        self._add_inline_runs(
            paragraph,
            text,
            font_color=_ORANGE,
            font_size=12,
            font_bold=True,
            font_name=_HEADING_FONT,
            east_asia_font=_HEADING_FONT,
        )

    @staticmethod
    def _add_inline_runs(
        paragraph: Any,
        text: str,
        *,
        font_size: Optional[float] = None,
        font_color: Optional[RGBColor] = None,
        font_bold: Optional[bool] = None,
        font_italic: Optional[bool] = None,
        font_name: Optional[str] = None,
        east_asia_font: Optional[str] = None,
    ) -> None:
        first_new_run = len(paragraph.runs)
        cursor = 0
        for match in _INLINE_PATTERN.finditer(text):
            if match.start() > cursor:
                run = paragraph.add_run(text[cursor : match.start()])

            token = match.group(0)
            if token.startswith(("**", "__")):
                content = token[2:-2]
                run = paragraph.add_run(content)
                run.bold = True
            else:
                content = token[1:-1]
                run = paragraph.add_run(content)
                run.italic = True
            cursor = match.end()

        if cursor < len(text):
            paragraph.add_run(text[cursor:])

        for run in paragraph.runs[first_new_run:]:
            if font_size is not None:
                run.font.size = Pt(font_size)
            if font_color is not None:
                run.font.color.rgb = font_color
            if font_name is not None or east_asia_font is not None:
                selected_font = font_name or _BODY_FONT
                selected_east_asia_font = east_asia_font or selected_font
                run.font.name = selected_font
                _set_font_mapping(run, selected_font, selected_east_asia_font)
            if font_bold is True:
                run.bold = True
            if font_italic is True:
                run.italic = True

    def _resolve_image_path(self, reference: str) -> Optional[Path]:
        reference = unquote(reference.strip())
        if reference.startswith("<") and reference.endswith(">"):
            reference = reference[1:-1]
        reference = reference.strip().strip('"').strip("'")
        if not reference:
            return None

        lookup_key = Path(reference).name
        candidate = self._figure_paths.get(reference) or self._figure_paths.get(lookup_key)
        if candidate is None:
            candidate = Path(reference)
            if not candidate.is_absolute():
                candidate = self.session_output_dir / candidate

        try:
            resolved = candidate.resolve()
        except OSError:
            return None

        if not _is_within(resolved, self.session_output_dir) or not resolved.is_file():
            return None
        return resolved

    def _add_image(self, alt_text: str, reference: str) -> None:
        image_path = self._resolve_image_path(reference)
        if image_path is None:
            paragraph = self.document.add_paragraph(style="Caption")
            label = alt_text.strip() or reference.strip()
            run = paragraph.add_run(f"[图片不可用: {label}]")
            _set_run_font(run, 9, _MUTED)
            run.italic = True
            return

        try:
            paragraph = self.document.add_paragraph(style="Figure")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(6.0))
            self._add_picture_effects(run)
            if alt_text.strip():
                caption = self.document.add_paragraph(style="Caption")
                caption.add_run(alt_text.strip())
        except Exception:
            paragraph = self.document.add_paragraph(style="Caption")
            label = alt_text.strip() or reference.strip()
            run = paragraph.add_run(f"[图片不可用: {label}]")
            _set_run_font(run, 9, _MUTED)
            run.italic = True

    @staticmethod
    def _add_picture_effects(run: Any) -> None:
        """Add a restrained border and outer shadow to an inline picture."""
        drawing = run._r.find(qn("w:drawing"))
        if drawing is None:
            return
        inline = drawing.find(qn("wp:inline"))
        if inline is None:
            return
        graphic = inline.find(qn("a:graphic"))
        graphic_data = graphic.find(qn("a:graphicData")) if graphic is not None else None
        picture = graphic_data.find(qn("pic:pic")) if graphic_data is not None else None
        if picture is None:
            return
        shape_properties = picture.find(qn("pic:spPr"))
        if shape_properties is None:
            return

        existing_line = shape_properties.find(qn("a:ln"))
        if existing_line is not None:
            shape_properties.remove(existing_line)

        line = OxmlElement("a:ln")
        line.set("w", "12700")
        solid_fill = OxmlElement("a:solidFill")
        line_color = OxmlElement("a:srgbClr")
        line_color.set("val", "D9D9D9")
        solid_fill.append(line_color)
        line.append(solid_fill)
        preset_dash = OxmlElement("a:prstDash")
        preset_dash.set("val", "solid")
        line.append(preset_dash)
        shape_properties.append(line)

        effect_list = shape_properties.find(qn("a:effectLst"))
        if effect_list is None:
            effect_list = OxmlElement("a:effectLst")
            shape_properties.append(effect_list)
        existing_shadow = effect_list.find(qn("a:outerShdw"))
        if existing_shadow is not None:
            effect_list.remove(existing_shadow)

        shadow = OxmlElement("a:outerShdw")
        shadow.set("blurRad", "63500")
        shadow.set("dist", "38100")
        shadow.set("dir", "5400000")
        shadow.set("rotWithShape", "0")
        shadow_color = OxmlElement("a:srgbClr")
        shadow_color.set("val", "808080")
        shadow_alpha = OxmlElement("a:alpha")
        shadow_alpha.set("val", "25000")
        shadow_color.append(shadow_alpha)
        shadow.append(shadow_color)
        effect_list.append(shadow)


def generate_word_report(
    markdown_content: str,
    output_path: PathLike,
    session_output_dir: PathLike,
    figures: Optional[Iterable[Dict[str, Any]]] = None,
) -> str:
    """Generate a DOCX report from Markdown and return its output path."""
    generator = WordReportGenerator(session_output_dir=session_output_dir, figures=figures)
    return generator.generate(markdown_content=markdown_content, output_path=output_path)


class WordReportRenderer:
    """Adapt the legacy Word generator to the report renderer interface."""

    format = ReportFormat.DOCX

    def __init__(
        self,
        generator: Callable[..., str] | None = None,
    ) -> None:
        self.generator = generate_word_report if generator is None else generator

    def render(
        self, *, markdown: str, document: ReportDocument, output_path: PathLike
    ) -> str:
        figures = [
            {
                "filename": chart.filename,
                "file_path": chart.file_path,
                "title": chart.title,
                "description": chart.description,
            }
            for chart in document.chart_artifacts
        ]
        return self.generator(
            markdown_content=markdown,
            output_path=output_path,
            session_output_dir=document.output_root,
            figures=figures,
        )
